"""Extract plain text from agent-uploaded documents (txt / md / pdf / docx / common text)."""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from xml.etree import ElementTree as ET

from app.core.exceptions import BusinessError

MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
# The upload byte limit already bounds request size. Never silently truncate
# extracted text: the Agent must either receive the entire document or get a
# clear error instead of planning tasks from a partial source.
MAX_EXTRACTED_CHARS = 1_000_000

SUPPORTED_HINT = "TXT / MD / PDF / DOCX / CSV / JSON / HTML / RTF"


@dataclass
class ExtractedDocument:
    name: str
    mime_type: str
    text: str
    char_count: int


def _decode_text_bytes(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf_text(content: bytes) -> str:
    # Prefer pypdf when available; fall back to a lightweight stream scanner.
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(BytesIO(content))
        parts: list[str] = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
        if text:
            return text
    except Exception:
        pass

    # Minimal PDF text draw operator scrape for simple text PDFs.
    chunks: list[str] = []
    for match in re.finditer(rb"\((?:\\.|[^\\)])*\)\s*Tj", content):
        raw = match.group(0)
        inner = raw[1 : raw.rfind(b")")]
        try:
            piece = inner.decode("utf-8", errors="ignore")
        except Exception:
            continue
        piece = piece.replace("\\n", "\n").replace("\\r", "").replace("\\t", "\t")
        piece = re.sub(r"\\(.)", r"\1", piece)
        if piece.strip():
            chunks.append(piece)
    for match in re.finditer(rb"\[(.*?)\]\s*TJ", content, flags=re.S):
        body = match.group(1)
        for part in re.finditer(rb"\((?:\\.|[^\\)])*\)", body):
            inner = part.group(0)[1:-1]
            try:
                piece = inner.decode("utf-8", errors="ignore")
            except Exception:
                continue
            piece = re.sub(r"\\(.)", r"\1", piece)
            if piece.strip():
                chunks.append(piece)
    text = "\n".join(chunks).strip()
    if not text:
        raise BusinessError("未能从 PDF 中提取到有效文本，请改用 TXT/DOCX 或复制正文", code=40058)
    return text


def _extract_docx_text(content: bytes) -> str:
    # Prefer python-docx when available.
    try:
        from docx import Document  # type: ignore

        document = Document(BytesIO(content))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            value = (paragraph.text or "").strip()
            if value:
                parts.append(value)
        for table in document.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                cells = [cell for cell in cells if cell]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if text:
            return text
    except Exception:
        pass

    # Fallback: OOXML is a zip; scrape w:t nodes from word/document.xml (+ headers/footers).
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
            targets = ["word/document.xml"]
            targets.extend(sorted(n for n in names if n.startswith("word/header") and n.endswith(".xml")))
            targets.extend(sorted(n for n in names if n.startswith("word/footer") and n.endswith(".xml")))
            chunks: list[str] = []
            for name in targets:
                if name not in names:
                    continue
                xml_bytes = archive.read(name)
                root = ET.fromstring(xml_bytes)
                # Collect text in document order.
                for node in root.iter():
                    tag = node.tag.rsplit("}", 1)[-1]
                    if tag == "t" and node.text:
                        chunks.append(node.text)
                    elif tag in {"p", "br", "cr", "tab"}:
                        chunks.append("\n" if tag != "tab" else "\t")
            text = "".join(chunks)
            text = re.sub(r"[ \t]+\n", "\n", text)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            if text:
                return text
    except zipfile.BadZipFile as exc:
        raise BusinessError("DOCX 文件损坏或不是有效 Word 文档", code=40059) from exc
    except Exception as exc:
        raise BusinessError("未能从 DOCX 中提取文本，请另存为 TXT/PDF 后重试", code=40059) from exc

    raise BusinessError("DOCX 中没有可识别文本", code=40057)


def _extract_html_text(content: bytes) -> str:
    raw = _decode_text_bytes(content)
    # Strip scripts/styles then tags.
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?is)<!--.*?-->", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</p\s*>", "\n", raw)
    raw = re.sub(r"(?i)</div\s*>", "\n", raw)
    raw = re.sub(r"(?i)</h[1-6]\s*>", "\n", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    raw = re.sub(r"&nbsp;", " ", raw)
    raw = re.sub(r"&amp;", "&", raw)
    raw = re.sub(r"&lt;", "<", raw)
    raw = re.sub(r"&gt;", ">", raw)
    raw = re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), raw)
    return _normalize_text(raw)


def _extract_rtf_text(content: bytes) -> str:
    raw = _decode_text_bytes(content)
    # Very small RTF control-word stripper for plain prose exports.
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    raw = re.sub(r"\\[{}\\]", lambda m: m.group(0)[-1], raw)
    return _normalize_text(raw)


def extract_document(*, filename: str, content: bytes, content_type: str | None = None) -> ExtractedDocument:
    if not content:
        raise BusinessError("文件内容为空", code=40054)
    if len(content) > MAX_DOCUMENT_BYTES:
        raise BusinessError("文档不能超过 8MB", code=40055)

    name = (filename or "document").strip() or "document"
    lower = name.lower()
    mime = (content_type or "").split(";")[0].strip().lower()

    if lower.endswith(".doc") and not lower.endswith(".docx"):
        raise BusinessError("暂不支持旧版 .doc，请用 Word「另存为」DOCX / PDF / TXT 后再上传", code=40056)

    if lower.endswith(".txt") or mime == "text/plain":
        text = _decode_text_bytes(content)
        kind_mime = "text/plain"
    elif lower.endswith(".md") or lower.endswith(".markdown") or mime in {"text/markdown", "text/x-markdown"}:
        text = _decode_text_bytes(content)
        kind_mime = "text/markdown"
    elif lower.endswith(".csv") or mime in {"text/csv", "application/csv"}:
        text = _decode_text_bytes(content)
        kind_mime = "text/csv"
    elif lower.endswith(".json") or mime in {"application/json", "text/json"}:
        text = _decode_text_bytes(content)
        kind_mime = "application/json"
    elif lower.endswith((".html", ".htm")) or mime in {"text/html", "application/xhtml+xml"}:
        text = _extract_html_text(content)
        kind_mime = "text/html"
    elif lower.endswith(".rtf") or mime == "application/rtf" or mime == "text/rtf":
        text = _extract_rtf_text(content)
        kind_mime = "application/rtf"
    elif lower.endswith(".pdf") or mime == "application/pdf":
        text = _extract_pdf_text(content)
        kind_mime = "application/pdf"
    elif lower.endswith(".docx") or mime in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/docx",
    }:
        text = _extract_docx_text(content)
        kind_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        raise BusinessError(f"仅支持 {SUPPORTED_HINT} 文档", code=40056)

    text = _normalize_text(text)
    if not text:
        raise BusinessError("文档中没有可识别文本", code=40057)
    if len(text) > MAX_EXTRACTED_CHARS:
        raise BusinessError("文档正文超过 100 万字，无法完整解析，请拆分后上传", code=40055)

    return ExtractedDocument(name=name, mime_type=kind_mime, text=text, char_count=len(text))
